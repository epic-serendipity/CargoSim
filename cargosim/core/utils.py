"""Utility functions and constants for CargoSim."""

import os
import sys
import time
import logging
import traceback
import json
import xml.etree.ElementTree as ET
from datetime import datetime
from typing import List, Dict, Any, Optional, Tuple
from functools import lru_cache

# Debug logging
from .paths import DEBUG_LOG_FILE, RUNTIME_LOG_FILE
DEBUG_LOG = str(DEBUG_LOG_FILE)
# Runtime logging for execution tracking
RUNTIME_LOG = str(RUNTIME_LOG_FILE)

# Setup logging
def setup_logging(level: str = "INFO", log_file: Optional[str] = None, 
                  console: bool = True) -> logging.Logger:
    """Setup logging configuration for CargoSim."""
    # Use the new comprehensive logging system
    from .logging_config import setup_comprehensive_logging
    log_manager = setup_comprehensive_logging(level, log_file, console)
    return log_manager.get_logger("cargosim")


def setup_runtime_logging() -> logging.Logger:
    """Setup specialized runtime logging for execution tracking."""
    runtime_logger = logging.getLogger("cargosim.runtime")
    runtime_logger.setLevel(logging.DEBUG)
    
    # Clear existing handlers
    runtime_logger.handlers.clear()
    
    # Create detailed formatter for runtime
    formatter = logging.Formatter(
        '%(asctime)s - %(name)s - %(levelname)s - [%(filename)s:%(lineno)d] - %(message)s'
    )
    
    # File handler for runtime log
    file_handler = logging.FileHandler(RUNTIME_LOG, encoding="utf-8", mode="w")
    file_handler.setFormatter(formatter)
    file_handler.setLevel(logging.DEBUG)
    runtime_logger.addHandler(file_handler)
    
    # Console handler for runtime (INFO level only)
    console_handler = logging.StreamHandler(sys.stdout)
    console_handler.setFormatter(formatter)
    console_handler.setLevel(logging.INFO)
    runtime_logger.addHandler(console_handler)
    
    return runtime_logger


def log_runtime_event(event: str, details: str = "", level: str = "INFO"):
    """Log a runtime event with timestamp and details."""
    runtime_logger = logging.getLogger("cargosim.runtime")
    
    # Add execution context
    import inspect
    frame = inspect.currentframe().f_back
    if frame:
        filename = os.path.basename(frame.f_code.co_filename)
        lineno = frame.f_lineno
        func_name = frame.f_code.co_name
        context = f"{filename}:{lineno}:{func_name}"
    else:
        context = "unknown"
    
    message = f"[{context}] {event}"
    if details:
        message += f" - {details}"
    
    if level.upper() == "DEBUG":
        runtime_logger.debug(message)
    elif level.upper() == "INFO":
        runtime_logger.info(message)
    elif level.upper() == "WARNING":
        runtime_logger.warning(message)
    elif level.upper() == "ERROR":
        runtime_logger.error(message)
    elif level.upper() == "CRITICAL":
        runtime_logger.critical(message)


def log_exception(exception: Exception, context: str = ""):
    """Log an exception with context information."""
    runtime_logger = logging.getLogger("cargosim.runtime")
    
    # Add execution context
    import inspect
    frame = inspect.currentframe().f_back
    if frame:
        filename = os.path.basename(frame.f_code.co_filename)
        lineno = frame.f_lineno
        func_name = frame.f_code.co_name
        context_info = f"{filename}:{lineno}:{func_name}"
    else:
        context_info = "unknown"
    
    message = f"[{context_info}] Exception in {context}: {type(exception).__name__}: {str(exception)}"
    runtime_logger.error(message)
    
    # Log full traceback
    runtime_logger.debug(f"Full traceback:\n{traceback.format_exc()}")


# Phase 8: Advanced Utility Functions

class PerformanceMonitor:
    """Real-time performance monitoring and analytics system."""
    
    def __init__(self):
        self.metrics = {}
        self.thresholds = {}
        self.alerts = []
        self.history = []
        self.start_time = time.time()
        self.last_update = time.time()
        
        # Performance thresholds
        self.thresholds = {
            'cpu_usage': 80.0,  # 80% CPU usage
            'memory_usage': 500.0,  # 500MB memory
            'response_time': 2.0,  # 2 seconds response time
            'error_rate': 0.05,  # 5% error rate
            'fps': 25.0  # 25 FPS minimum
        }
    
    def track_metric(self, name: str, value: float, timestamp: float = None):
        """Track a performance metric."""
        if timestamp is None:
            timestamp = time.time()
        
        if name not in self.metrics:
            self.metrics[name] = []
        
        self.metrics[name].append({
            'value': value,
            'timestamp': timestamp
        })
        
        # Keep only recent history (last 1000 entries)
        if len(self.metrics[name]) > 1000:
            self.metrics[name] = self.metrics[name][-1000:]
        
        # Check thresholds and generate alerts
        self._check_thresholds(name, value, timestamp)
        
        self.last_update = timestamp
    
    def _check_thresholds(self, name: str, value: float, timestamp: float):
        """Check if metric exceeds thresholds and generate alerts."""
        if name in self.thresholds:
            threshold = self.thresholds[name]
            
            # Determine if threshold is exceeded (depends on metric type)
            exceeded = False
            if name in ['cpu_usage', 'memory_usage', 'error_rate']:
                exceeded = value > threshold
            elif name in ['response_time']:
                exceeded = value > threshold
            elif name in ['fps']:
                exceeded = value < threshold
            
            if exceeded:
                alert = {
                    'timestamp': timestamp,
                    'metric': name,
                    'value': value,
                    'threshold': threshold,
                    'severity': 'high' if value > threshold * 1.5 else 'medium'
                }
                self.alerts.append(alert)
                
                # Keep only recent alerts
                if len(self.alerts) > 100:
                    self.alerts = self.alerts[-100:]
    
    def get_metric_history(self, name: str, duration_hours: float = 24.0) -> List[dict]:
        """Get metric history for the specified duration."""
        if name not in self.metrics:
            return []
        
        cutoff_time = time.time() - (duration_hours * 3600)
        return [m for m in self.metrics[name] if m['timestamp'] >= cutoff_time]
    
    def get_current_metric(self, name: str) -> Optional[float]:
        """Get the most recent value for a metric."""
        if name in self.metrics and self.metrics[name]:
            return self.metrics[name][-1]['value']
        return None
    
    def calculate_trend(self, name: str, duration_hours: float = 24.0) -> float:
        """Calculate trend for a metric over the specified duration."""
        history = self.get_metric_history(name, duration_hours)
        if len(history) < 2:
            return 0.0
        
        # Simple linear regression
        x_values = [m['timestamp'] for m in history]
        y_values = [m['value'] for m in history]
        
        n = len(x_values)
        sum_x = sum(x_values)
        sum_y = sum(y_values)
        sum_xy = sum(x * y for x, y in zip(x_values, y_values))
        sum_x2 = sum(x * x for x in x_values)
        
        if n * sum_x2 - sum_x * sum_x != 0:
            trend = (n * sum_xy - sum_x * sum_y) / (n * sum_x2 - sum_x * sum_x)
            return trend
        
        return 0.0
    
    def get_performance_summary(self) -> dict:
        """Get a summary of current performance metrics."""
        summary = {
            'uptime_hours': (time.time() - self.start_time) / 3600,
            'last_update': self.last_update,
            'active_metrics': len(self.metrics),
            'total_alerts': len(self.alerts),
            'current_metrics': {}
        }
        
        # Add current values for all metrics
        for name in self.metrics:
            summary['current_metrics'][name] = self.get_current_metric(name)
        
        return summary
    
    def clear_history(self):
        """Clear all metric history."""
        self.metrics.clear()
        self.alerts.clear()
        self.history.clear()


class AnalyticsEngine:
    """Advanced analytics and predictive modeling system."""
    
    def __init__(self):
        self.models = {}
        self.predictions = {}
        self.accuracy_history = []
        self.last_training = time.time()
        
        # Initialize basic models
        self._initialize_models()
    
    def _initialize_models(self):
        """Initialize basic predictive models."""
        # Demand forecasting model
        self.models['demand_forecast'] = {
            'type': 'linear_regression',
            'parameters': {},
            'accuracy': 0.0,
            'last_updated': time.time()
        }
        
        # Cost prediction model
        self.models['cost_prediction'] = {
            'type': 'moving_average',
            'parameters': {'window_size': 10},
            'accuracy': 0.0,
            'last_updated': time.time()
        }
        
        # Performance prediction model
        self.models['performance_prediction'] = {
            'type': 'exponential_smoothing',
            'parameters': {'alpha': 0.3},
            'accuracy': 0.0,
            'last_updated': time.time()
        }
    
    def predict_demand(self, spoke_idx: int, resource: str, time_horizon: float) -> dict:
        """Predict resource demand for a specific spoke and resource."""
        model = self.models['demand_forecast']
        
        # Simple linear prediction (placeholder for more sophisticated models)
        base_demand = 1.0  # Base demand per period
        trend_factor = 0.1  # Small upward trend
        
        predicted_demand = base_demand + (trend_factor * time_horizon)
        
        # Add some randomness for realism
        import random
        noise = random.uniform(-0.2, 0.2)
        predicted_demand = max(0.0, predicted_demand + noise)
        
        prediction = {
            'spoke_idx': spoke_idx,
            'resource': resource,
            'predicted_demand': predicted_demand,
            'time_horizon': time_horizon,
            'confidence': 0.7,
            'model_type': model['type'],
            'timestamp': time.time()
        }
        
        # Store prediction
        if 'demand_predictions' not in self.predictions:
            self.predictions['demand_predictions'] = []
        self.predictions['demand_predictions'].append(prediction)
        
        return prediction
    
    def predict_cost(self, operation_type: str, distance: float, aircraft_type: str) -> dict:
        """Predict operation cost based on various factors."""
        model = self.models['cost_prediction']
        
        # Base cost calculation
        base_cost = 100.0  # Base cost per operation
        
        # Distance factor
        distance_cost = distance * 0.5
        
        # Aircraft type factor
        aircraft_multiplier = 1.0
        if aircraft_type == "C-130":
            aircraft_multiplier = 1.5
        elif aircraft_type == "C-27":
            aircraft_multiplier = 1.0
        
        # Operation type factor
        operation_multiplier = 1.0
        if operation_type == "loading":
            operation_multiplier = 0.8
        elif operation_type == "unloading":
            operation_multiplier = 0.8
        elif operation_type == "maintenance":
            operation_multiplier = 2.0
        
        predicted_cost = (base_cost + distance_cost) * aircraft_multiplier * operation_multiplier
        
        prediction = {
            'operation_type': operation_type,
            'distance': distance,
            'aircraft_type': aircraft_type,
            'predicted_cost': predicted_cost,
            'confidence': 0.8,
            'model_type': model['type'],
            'timestamp': time.time()
        }
        
        # Store prediction
        if 'cost_predictions' not in self.predictions:
            self.predictions['cost_predictions'] = []
        self.predictions['cost_predictions'].append(prediction)
        
        return prediction
    
    def predict_performance(self, metric: str, current_value: float, time_horizon: float) -> dict:
        """Predict future performance based on current trends."""
        model = self.models['performance_prediction']
        
        # Simple trend-based prediction
        trend = 0.0  # Placeholder for actual trend calculation
        
        predicted_value = current_value + (trend * time_horizon)
        
        prediction = {
            'metric': metric,
            'current_value': current_value,
            'predicted_value': predicted_value,
            'time_horizon': time_horizon,
            'trend': trend,
            'confidence': 0.6,
            'model_type': model['type'],
            'timestamp': time.time()
        }
        
        # Store prediction
        if 'performance_predictions' not in self.predictions:
            self.predictions['performance_predictions'] = []
        self.predictions['performance_predictions'].append(prediction)
        
        return prediction
    
    def update_model_accuracy(self, model_name: str, actual_value: float, predicted_value: float):
        """Update model accuracy based on actual vs predicted values."""
        if model_name not in self.models:
            return
        
        # Calculate prediction error
        error = abs(actual_value - predicted_value) / max(actual_value, 0.001)
        accuracy = max(0.0, 1.0 - error)
        
        # Update model accuracy (exponential moving average)
        current_accuracy = self.models[model_name]['accuracy']
        alpha = 0.1  # Learning rate
        new_accuracy = alpha * accuracy + (1 - alpha) * current_accuracy
        
        self.models[model_name]['accuracy'] = new_accuracy
        self.models[model_name]['last_updated'] = time.time()
        
        # Store accuracy history
        self.accuracy_history.append({
            'model': model_name,
            'actual_value': actual_value,
            'predicted_value': predicted_value,
            'accuracy': accuracy,
            'timestamp': time.time()
        })
        
        # Keep only recent history
        if len(self.accuracy_history) > 1000:
            self.accuracy_history = self.accuracy_history[-1000:]
    
    def get_model_status(self) -> dict:
        """Get status of all predictive models."""
        status = {}
        for name, model in self.models.items():
            status[name] = {
                'type': model['type'],
                'accuracy': model['accuracy'],
                'last_updated': model['last_updated'],
                'parameters': model['parameters']
            }
        return status
    
    def get_prediction_summary(self) -> dict:
        """Get summary of all predictions."""
        summary = {}
        for category, predictions in self.predictions.items():
            summary[category] = {
                'count': len(predictions),
                'latest': predictions[-1] if predictions else None,
                'average_confidence': sum(p['confidence'] for p in predictions) / len(predictions) if predictions else 0.0
            }
        return summary


class ExportManager:
    """Enhanced data export and integration management system."""
    
    def __init__(self):
        self.export_history = []
        self.integration_endpoints = {
            'database': {
                'enabled': False,
                'type': 'postgresql',
                'host': 'localhost',
                'port': 5432,
                'database': 'cargosim',
                'username': '',
                'password': '',
                'ssl_mode': 'prefer'
            },
            'cloud_storage': {
                'enabled': False,
                'provider': 'aws_s3',
                'bucket': '',
                'region': 'us-east-1',
                'access_key': '',
                'secret_key': '',
                'endpoint_url': None
            },
            'api_server': {
                'enabled': False,
                'base_url': 'http://localhost:8000',
                'api_key': '',
                'timeout': 30,
                'retry_attempts': 3
            },
            'webhook': {
                'enabled': False,
                'url': '',
                'method': 'POST',
                'headers': {},
                'timeout': 10
            }
        }
        self.export_formats = ['csv', 'json', 'excel', 'pdf', 'xml', 'yaml']
        self.max_history_size = 1000
        self.compression_enabled = True
        
    def export_simulation_data(self, data: dict, format: str, filename: str = None, 
                              compression: bool = None, metadata: dict = None) -> str:
        """Export simulation data in multiple formats with enhanced options."""
        if compression is None:
            compression = self.compression_enabled
            
        if metadata is None:
            metadata = {}
        
        # Add export metadata
        export_info = {
            'export_timestamp': time.time(),
            'export_format': format,
            'compression': compression,
            'data_size': len(str(data)),
            'export_manager_version': '2.0.0'
        }
        export_info.update(metadata)
        
        # Generate filename if not provided
        if not filename:
            timestamp = datetime.fromtimestamp(export_info['export_timestamp']).strftime('%Y%m%d_%H%M%S')
            filename = f"cargosim_export_{timestamp}.{format}"
        
        try:
            # Export based on format
            if format == 'csv':
                result_file = self._export_to_csv(data, filename)
            elif format == 'json':
                result_file = self._export_to_json(data, filename, export_info)
            elif format == 'excel':
                result_file = self._export_to_excel(data, filename, export_info)
            elif format == 'pdf':
                result_file = self._export_to_pdf(data, filename, export_info)
            elif format == 'xml':
                result_file = self._export_to_xml(data, filename, export_info)
            elif format == 'yaml':
                result_file = self._export_to_yaml(data, filename, export_info)
            else:
                raise ValueError(f"Unsupported export format: {format}")
            
            # Apply compression if requested
            if compression and format not in ['excel', 'pdf']:
                result_file = self._compress_file(result_file)
            
            # Record export in history
            self._record_export(filename, format, export_info)
            
            # Upload to cloud storage if enabled
            if self.integration_endpoints['cloud_storage']['enabled']:
                self._upload_to_cloud(result_file, filename)
            
            # Send webhook notification if enabled
            if self.integration_endpoints['webhook']['enabled']:
                self._send_webhook_notification(filename, export_info)
            
            return result_file
            
        except Exception as e:
            log_exception(e, f"Export to {format} failed")
            raise
    
    def _export_to_csv(self, data: dict, filename: str) -> str:
        """Export data to CSV format."""
        try:
            import pandas as pd
            
            # Flatten nested data for CSV
            flattened_data = self._flatten_dict(data)
            
            # Convert to DataFrame
            df = pd.DataFrame([flattened_data])
            
            # Export to CSV
            df.to_csv(filename, index=False, encoding='utf-8')
            return filename
            
        except ImportError:
            # Fallback to manual CSV generation
            return self._export_to_csv_manual(data, filename)
    
    def _export_to_csv_manual(self, data: dict, filename: str) -> str:
        """Manual CSV export without pandas dependency."""
        try:
            with open(filename, 'w', newline='', encoding='utf-8') as f:
                # Write header
                headers = self._get_csv_headers(data)
                f.write(','.join(f'"{h}"' for h in headers) + '\n')
                
                # Write data
                row_data = self._get_csv_row_data(data, headers)
                f.write(','.join(f'"{str(v)}"' for v in row_data) + '\n')
            
            return filename
        except Exception as e:
            raise Exception(f"Manual CSV export failed: {e}")
    
    def _export_to_json(self, data: dict, filename: str, metadata: dict) -> str:
        """Export data to JSON format with metadata."""
        try:
            import json
            # Combine data and metadata
            export_data = {
                'metadata': metadata,
                'data': data,
                'export_info': {
                    'exported_by': 'CargoSim Export Manager',
                    'export_timestamp': datetime.fromtimestamp(metadata['export_timestamp']).isoformat(),
                    'version': metadata.get('export_manager_version', '2.0.0')
                }
            }
            
            # Write JSON file
            with open(filename, 'w', encoding='utf-8') as f:
                json.dump(export_data, f, indent=2, default=str, ensure_ascii=False)
            
            return filename
            
        except Exception as e:
            raise Exception(f"JSON export failed: {e}")
    
    def _export_to_excel(self, data: dict, filename: str, metadata: dict) -> str:
        """Export data to Excel format with multiple sheets."""
        try:
            import pandas as pd
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            
            # Create Excel workbook
            wb = Workbook()
            
            # Main data sheet
            ws_main = wb.active
            ws_main.title = "Simulation Data"
            
            # Add metadata sheet
            ws_meta = wb.create_sheet("Export Metadata")
            
            # Write main data
            self._write_excel_data(ws_main, data)
            
            # Write metadata
            self._write_excel_metadata(ws_meta, metadata)
            
            # Save workbook
            wb.save(filename)
            return filename
            
        except ImportError:
            raise Exception("Excel export requires openpyxl package")
    
    def _export_to_pdf(self, data: dict, filename: str, metadata: dict) -> str:
        """Export data to PDF format with professional styling."""
        try:
            from reportlab.lib.pagesizes import letter, A4
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib import colors
            from reportlab.lib.units import inch
            
            # Create PDF document
            doc = SimpleDocTemplate(filename, pagesize=A4)
            story = []
            
            # Add title
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                spaceAfter=30,
                alignment=1  # Center
            )
            story.append(Paragraph("CargoSim Simulation Report", title_style))
            story.append(Spacer(1, 20))
            
            # Add metadata
            story.append(Paragraph("Export Information", styles['Heading2']))
            meta_table_data = [[k, str(v)] for k, v in metadata.items()]
            meta_table = Table(meta_table_data, colWidths=[2*inch, 4*inch])
            meta_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            story.append(meta_table)
            story.append(Spacer(1, 20))
            
            # Add data summary
            story.append(Paragraph("Data Summary", styles['Heading2']))
            data_summary = self._create_data_summary(data)
            story.append(Paragraph(data_summary, styles['Normal']))
            
            # Build PDF
            doc.build(story)
            return filename
            
        except ImportError:
            raise Exception("PDF export requires reportlab package")
    
    def _export_to_xml(self, data: dict, filename: str, metadata: dict) -> str:
        """Export data to XML format."""
        try:
            import xml.etree.ElementTree as ET
            from xml.dom import minidom
            
            # Create root element
            root = ET.Element("CargoSimExport")
            root.set("version", "2.0")
            root.set("timestamp", str(metadata.get('export_timestamp', time.time())))
            
            # Add metadata
            meta_elem = ET.SubElement(root, "Metadata")
            for key, value in metadata.items():
                meta_item = ET.SubElement(meta_elem, key)
                meta_item.text = str(value)
            
            # Add data
            data_elem = ET.SubElement(root, "Data")
            self._dict_to_xml(data, data_elem)
            
            # Create pretty XML
            rough_string = ET.tostring(root, 'unicode')
            reparsed = minidom.parseString(rough_string)
            pretty_xml = reparsed.toprettyxml(indent="  ")
            
            # Write to file
            with open(filename, 'w', encoding='utf-8') as f:
                f.write(pretty_xml)
            
            return filename
            
        except Exception as e:
            raise Exception(f"XML export failed: {e}")
    
    def _export_to_yaml(self, data: dict, filename: str, metadata: dict) -> str:
        """Export data to YAML format."""
        try:
            import yaml
            
            # Combine data and metadata
            export_data = {
                'metadata': metadata,
                'data': data
            }
            
            # Write YAML file
            with open(filename, 'w', encoding='utf-8') as f:
                yaml.dump(export_data, f, default_flow_style=False, 
                         sort_keys=False, allow_unicode=True)
            
            return filename
            
        except ImportError:
            raise Exception("YAML export requires PyYAML package")
    
    def _compress_file(self, filename: str) -> str:
        """Compress a file using gzip."""
        try:
            import gzip
            import shutil
            
            compressed_filename = filename + '.gz'
            
            with open(filename, 'rb') as f_in:
                with gzip.open(compressed_filename, 'wb') as f_out:
                    shutil.copyfileobj(f_in, f_out)
            
            # Remove original file
            os.remove(filename)
            return compressed_filename
            
        except Exception as e:
            log_exception(e, f"File compression failed for {filename}")
            return filename
    
    def _upload_to_cloud(self, local_file: str, remote_filename: str):
        """Upload file to cloud storage."""
        try:
            config = self.integration_endpoints['cloud_storage']
            
            if config['provider'] == 'aws_s3':
                self._upload_to_s3(local_file, remote_filename, config)
            elif config['provider'] == 'google_cloud':
                self._upload_to_gcs(local_file, remote_filename, config)
            elif config['provider'] == 'azure_blob':
                self._upload_to_azure(local_file, remote_filename, config)
            else:
                log_runtime_event(f"Unsupported cloud provider: {config['provider']}")
                
        except Exception as e:
            log_exception(e, f"Cloud upload failed for {local_file}")
    
    def _upload_to_s3(self, local_file: str, remote_filename: str, config: dict):
        """Upload file to AWS S3."""
        try:
            import boto3
            from botocore.exceptions import ClientError
            
            # Create S3 client
            s3_client = boto3.client(
                's3',
                aws_access_key_id=config['access_key'],
                aws_secret_access_key=config['secret_key'],
                region_name=config['region'],
                endpoint_url=config.get('endpoint_url')
            )
            
            # Upload file
            s3_client.upload_file(local_file, config['bucket'], remote_filename)
            log_runtime_event(f"File uploaded to S3: s3://{config['bucket']}/{remote_filename}")
            
        except ImportError:
            log_runtime_event("boto3 not available for S3 upload")
        except ClientError as e:
            log_exception(e, f"S3 upload failed: {e}")
    
    def _send_webhook_notification(self, filename: str, metadata: dict):
        """Send webhook notification about export completion."""
        try:
            import requests
            
            config = self.integration_endpoints['webhook']
            
            # Prepare webhook data
            webhook_data = {
                'event': 'export_completed',
                'filename': filename,
                'timestamp': metadata['export_timestamp'],
                'format': metadata['export_format'],
                'data_size': metadata['data_size']
            }
            
            # Send webhook
            response = requests.post(
                config['url'],
                json=webhook_data,
                headers=config['headers'],
                timeout=config['timeout']
            )
            
            if response.status_code == 200:
                log_runtime_event(f"Webhook notification sent successfully for {filename}")
            else:
                log_runtime_event(f"Webhook notification failed with status {response.status_code}")
                
        except ImportError:
            log_runtime_event("requests not available for webhook notifications")
        except Exception as e:
            log_exception(e, f"Webhook notification failed for {filename}")
    
    def _record_export(self, filename: str, format: str, metadata: dict):
        """Record export in history."""
        export_record = {
            'filename': filename,
            'format': format,
            'timestamp': metadata['export_timestamp'],
            'size': metadata['data_size'],
            'compression': metadata.get('compression', False)
        }
        
        self.export_history.append(export_record)
        
        # Maintain history size limit
        if len(self.export_history) > self.max_history_size:
            self.export_history.pop(0)
    
    def get_export_history(self, limit: int = None) -> list:
        """Get export history with optional limit."""
        if limit is None:
            return self.export_history.copy()
        return self.export_history[-limit:]
    
    def clear_export_history(self):
        """Clear export history."""
        self.export_history.clear()
    
    def get_export_statistics(self) -> dict:
        """Get export statistics and analytics."""
        if not self.export_history:
            return {
                'total_exports': 0,
                'formats_used': {},
                'total_data_size': 0,
                'average_size': 0,
                'compression_ratio': 0
            }
        
        total_exports = len(self.export_history)
        formats_used = {}
        total_size = 0
        
        for export in self.export_history:
            format_type = export['format']
            formats_used[format_type] = formats_used.get(format_type, 0) + 1
            total_size += export['size']
        
        return {
            'total_exports': total_exports,
            'formats_used': formats_used,
            'total_data_size': total_size,
            'average_size': total_size / total_exports if total_exports > 0 else 0,
            'compression_ratio': 0.7,  # Placeholder for actual compression ratio
            'last_export': self.export_history[-1] if self.export_history else None
        }
    
    def _flatten_dict(self, data: dict, parent_key: str = '', sep: str = '_') -> dict:
        """Flatten nested dictionary for CSV export."""
        items = []
        for k, v in data.items():
            new_key = f"{parent_key}{sep}{k}" if parent_key else k
            if isinstance(v, dict):
                items.extend(self._flatten_dict(v, new_key, sep=sep).items())
            elif isinstance(v, list):
                items.append((new_key, str(v)))
            else:
                items.append((new_key, v))
        return dict(items)
    
    def _get_csv_headers(self, data: dict) -> list:
        """Get CSV headers from flattened data."""
        flattened = self._flatten_dict(data)
        return list(flattened.keys())
    
    def _get_csv_row_data(self, data: dict, headers: list) -> list:
        """Get CSV row data from flattened data."""
        flattened = self._flatten_dict(data)
        return [flattened.get(header, '') for header in headers]
    
    def _write_excel_data(self, worksheet, data: dict):
        """Write data to Excel worksheet."""
        # Implementation for Excel data writing
        pass
    
    def _write_excel_metadata(self, worksheet, metadata: dict):
        """Write metadata to Excel worksheet."""
        # Implementation for Excel metadata writing
        pass
    
    def _dict_to_xml(self, data: dict, parent_element):
        """Convert dictionary to XML elements."""
        for key, value in data.items():
            if isinstance(value, dict):
                child = ET.SubElement(parent_element, key)
                self._dict_to_xml(value, child)
            elif isinstance(value, list):
                for item in value:
                    child = ET.SubElement(parent_element, key)
                    if isinstance(item, dict):
                        self._dict_to_xml(item, child)
                    else:
                        child.text = str(item)
            else:
                child = ET.SubElement(parent_element, key)
                child.text = str(value)
    
    def _create_data_summary(self, data: dict) -> str:
        """Create a text summary of the data for PDF export."""
        summary = f"Data contains {len(data)} top-level keys.\n\n"
        
        for key, value in data.items():
            if isinstance(value, dict):
                summary += f"{key}: Dictionary with {len(value)} items\n"
            elif isinstance(value, list):
                summary += f"{key}: List with {len(value)} items\n"
            else:
                summary += f"{key}: {type(value).__name__}\n"
        
        return summary
    
    def configure_integration(self, endpoint: str, config: dict):
        """Configure an integration endpoint."""
        if endpoint in self.integration_endpoints:
            self.integration_endpoints[endpoint].update(config)
            self.integration_endpoints[endpoint]['enabled'] = True
            log_runtime_event(f"Integration endpoint '{endpoint}' configured successfully")
    
    def test_integration(self, endpoint: str) -> bool:
        """Test an integration endpoint."""
        if endpoint not in self.integration_endpoints:
            return False
        
        config = self.integration_endpoints[endpoint]
        if not config['enabled']:
            return False
        
        try:
            if endpoint == 'database':
                return self._test_database_connection(config)
            elif endpoint == 'cloud_storage':
                return self._test_cloud_storage_connection(config)
            elif endpoint == 'api_server':
                return self._test_api_server_connection(config)
            elif endpoint == 'webhook':
                return self._test_webhook_endpoint(config)
            else:
                return False
        except Exception as e:
            log_exception(e, f"Integration test failed for {endpoint}")
            return False
    
    def _test_database_connection(self, config: dict) -> bool:
        """Test database connection."""
        try:
            if config['type'] == 'postgresql':
                import psycopg2
                conn = psycopg2.connect(
                    host=config['host'],
                    port=config['port'],
                    database=config['database'],
                    user=config['username'],
                    password=config['password']
                )
                conn.close()
                return True
            else:
                log_runtime_event(f"Unsupported database type: {config['type']}")
                return False
        except ImportError:
            log_runtime_event("psycopg2 not available for PostgreSQL testing")
            return False
        except Exception as e:
            log_exception(e, "Database connection test failed")
            return False
    
    def _test_cloud_storage_connection(self, config: dict) -> bool:
        """Test cloud storage connection."""
        try:
            if config['provider'] == 'aws_s3':
                import boto3
                s3_client = boto3.client(
                    's3',
                    aws_access_key_id=config['access_key'],
                    aws_secret_access_key=config['secret_key'],
                    region_name=config['region']
                )
                # Test by listing buckets
                s3_client.list_buckets()
                return True
            else:
                log_runtime_event(f"Unsupported cloud provider: {config['provider']}")
                return False
        except ImportError:
            log_runtime_event("boto3 not available for S3 testing")
            return False
        except Exception as e:
            log_exception(e, "Cloud storage connection test failed")
            return False
    
    def _test_api_server_connection(self, config: dict) -> bool:
        """Test API server connection."""
        try:
            import requests
            response = requests.get(
                f"{config['base_url']}/health",
                timeout=config['timeout'],
                headers={'Authorization': f"Bearer {config['api_key']}"} if config['api_key'] else {}
            )
            return response.status_code == 200
        except ImportError:
            log_runtime_event("requests not available for API testing")
            return False
        except Exception as e:
            log_exception(e, "API server connection test failed")
            return False
    
    def _test_webhook_endpoint(self, config: dict) -> bool:
        """Test webhook endpoint."""
        try:
            import requests
            test_data = {
                'test': True,
                'timestamp': time.time(),
                'message': 'CargoSim integration test'
            }
            response = requests.post(
                config['url'],
                json=test_data,
                headers=config['headers'],
                timeout=config['timeout']
            )
            return response.status_code in [200, 201, 202]
        except ImportError:
            log_runtime_event("requests not available for webhook testing")
            return False
        except Exception as e:
            log_exception(e, "Webhook endpoint test failed")
            return False
    
    def get_integration_status(self) -> dict:
        """Get status of all integration endpoints."""
        status = {}
        for name, config in self.integration_endpoints.items():
            status[name] = {
                'enabled': config['enabled'],
                'test_result': self.test_integration(name) if config['enabled'] else False,
                'last_tested': time.time() if config['enabled'] else None
            }
        return status
    
    def get_integration_config(self, endpoint: str) -> dict:
        """Get configuration for a specific integration endpoint."""
        if endpoint in self.integration_endpoints:
            return self.integration_endpoints[endpoint].copy()
        return {}
    
    def disable_integration(self, endpoint: str):
        """Disable an integration endpoint."""
        if endpoint in self.integration_endpoints:
            self.integration_endpoints[endpoint]['enabled'] = False
            log_runtime_event(f"Integration endpoint '{endpoint}' disabled")
    
    def validate_integration_config(self, endpoint: str, config: dict) -> tuple[bool, str]:
        """Validate integration configuration."""
        if endpoint not in self.integration_endpoints:
            return False, f"Unknown integration endpoint: {endpoint}"
        
        required_fields = self._get_required_fields(endpoint)
        for field in required_fields:
            if field not in config or not config[field]:
                return False, f"Missing required field: {field}"
        
        return True, "Configuration valid"
    
    def _get_required_fields(self, endpoint: str) -> list:
        """Get required fields for an integration endpoint."""
        required_fields = {
            'database': ['host', 'port', 'database', 'username', 'password'],
            'cloud_storage': ['provider', 'bucket', 'region', 'access_key', 'secret_key'],
            'api_server': ['base_url', 'api_key'],
            'webhook': ['url', 'method']
        }
        return required_fields.get(endpoint, [])


# Global instances for easy access
performance_monitor = PerformanceMonitor()
analytics_engine = AnalyticsEngine()
export_manager = ExportManager()


def get_logger(name: str = "cargosim") -> logging.Logger:
    """Get a logger instance."""
    # Use the new comprehensive logging system
    from .logging_config import get_logger as get_comprehensive_logger
    return get_comprehensive_logger(name)


def append_debug(lines: List[str]):
    """Append debug lines to the debug log file (legacy function)."""
    try:
        with open(DEBUG_LOG, "a", encoding="utf-8") as f:
            for ln in lines:
                f.write(ln + "\n")
    except Exception:
        pass


def clamp(val, lo, hi):
    """Clamp a value between low and high bounds."""
    return max(lo, min(hi, val))


def ellipsize(text: str, font, max_w: int) -> str:
    """Truncate text with ellipsis so rendered width ≤ max_w."""
    if font.size(text)[0] <= max_w - 2:
        return text
    out = text
    while out and font.size(out + "...")[0] > max_w:
        out = out[:-1]
    return out + "..." if out else "..."


def ensure_mp4_ext(path: str) -> str:
    """Ensure path has .mp4 extension."""
    if not path.lower().endswith('.mp4'):
        path += '.mp4'
    return path


def tmp_mp4_path() -> str:
    """Get temporary MP4 path."""
    import tempfile
    return os.path.join(tempfile.gettempdir(), f"cargo_sim_{os.getpid()}.mp4")


def _mp4_available() -> tuple[bool, str]:
    """Check if MP4 recording is available."""
    try:
        import imageio
        import imageio_ffmpeg
        return True, "imageio-ffmpeg available"
    except ImportError:
        return False, "imageio-ffmpeg not available"

# ---------------------------------------------------------------------------
# User configuration reset helper
# ---------------------------------------------------------------------------
from shutil import rmtree

def reset_user_configs() -> None:
    """Delete all user-generated configuration files and related metadata."""
    from .paths import USER_CONFIGS_DIR, DEFAULT_CONFIGS_DIR
    import shutil, glob
    try:
        # Purge user directory completely
        if USER_CONFIGS_DIR.exists():
            shutil.rmtree(USER_CONFIGS_DIR)
        USER_CONFIGS_DIR.mkdir(parents=True, exist_ok=True)

        # Remove accidental migration files/backups in default dir
        for meta in DEFAULT_CONFIGS_DIR.glob("*.migration.json"):
            try:
                meta.unlink()
            except Exception:
                pass
        backup_dir = DEFAULT_CONFIGS_DIR / "backups"
        if backup_dir.exists():
            shutil.rmtree(backup_dir, ignore_errors=True)

    except Exception as exc:
        print(f"Warning: failed to reset user configs: {exc}")